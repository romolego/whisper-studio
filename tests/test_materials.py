from __future__ import annotations

import base64
import io
import json
import tempfile
import threading
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException

from backend import material_exports, materials, materials_api, storage

# Small real PNG; uploaded bytes must survive unchanged.
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1cAAAAASUVORK5CYII=")


class MaterialsTests(unittest.TestCase):
    def setUp(self):
        work = Path(__file__).resolve().parents[1] / "run"
        work.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=work)
        self.root = Path(self.temp.name)
        self.patches = [patch.object(storage, "DATA_DIR", self.root),
                        patch.object(storage, "DB_PATH", self.root / "test.db"),
                        patch.object(storage, "UPLOADS_DIR", self.root / "sessions"),
                        patch.dict(storage.DEFAULT_SETTINGS, {"managed_storage_root": str(self.root / "sessions")})]
        for item in self.patches:
            item.start()
        storage.init_db()
        self.folder = storage.create_folder("Материалы")
        self.project = storage.create_session("Тест материалов", self.folder["id"])["id"]

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def test_text_roundtrip_and_files_beside_project(self):
        text = "# Заголовок\n\n**Решили:** проверить.\n- Первый пункт\n  Пробелы сохранены.\n"
        saved = materials.save_text(self.project, "minutes", text, 0)
        self.assertEqual(saved["revision"], 1)
        self.assertEqual(materials.get(self.project)["minutes"], text)
        root = storage.session_dir(self.project) / "materials"
        self.assertEqual((root / "Протокол.txt").read_text(encoding="utf-8"), text)
        doc = Document(root / "Протокол.docx")
        self.assertEqual(doc.paragraphs[0].text, "Заголовок")
        self.assertIn("Решили: проверить.", [p.text for p in doc.paragraphs])
        self.assertTrue(any(run.bold for paragraph in doc.paragraphs for run in paragraph.runs))
        self.assertEqual(materials.text_file(self.project, "minutes", "txt")[1], text.encode("utf-8"))

    def test_stale_revision_cannot_overwrite(self):
        materials.save_text(self.project, "summary", "Первый результат", 0)
        with self.assertRaises(materials.Conflict):
            materials.save_text(self.project, "summary", "Устаревшее окно", 0)
        self.assertEqual(materials.get(self.project)["summary"], "Первый результат")

    def test_simultaneous_agent_and_browser_exactly_one_writer(self):
        def write(value):
            try:
                materials.save_text(self.project, "summary", value, 0)
                return "saved"
            except materials.Conflict:
                return "conflict"
        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertCountEqual(list(pool.map(write, ["Агент", "Браузер"])), ["saved", "conflict"])

    def test_clearing_text_removes_generated_files(self):
        materials.save_text(self.project, "summary", "Было", 0)
        data = materials.save_text(self.project, "summary", "  ", 1)
        self.assertFalse(data["summary"].strip())
        self.assertFalse((storage.session_dir(self.project) / "materials" / "Резюме.docx").exists())
        with self.assertRaises(ValueError):
            materials.text_file(self.project, "summary", "docx")

    def test_images_notes_order_and_delete(self):
        data = materials.add_image(self.project, "схема.png", PNG, 0)
        data = materials.add_image(self.project, "вторая.png", PNG, data["revision"])
        first, second = data["images"]
        path, name = materials.image_file(self.project, first["id"])
        self.assertEqual(path.read_bytes(), PNG)
        self.assertEqual(name, "схема.png")
        data = materials.edit_images(self.project, [{"id": second["id"], "note": "Лучший вариант"},
                                                    {"id": first["id"], "note": "Черновик"}], data["revision"])
        self.assertEqual(data["images"][0]["id"], second["id"])
        materials.edit_images(self.project, [{"id": second["id"], "note": "Лучший вариант"}], data["revision"])
        self.assertFalse(path.exists())
        with self.assertRaises(KeyError):
            materials.image_file(self.project, first["id"])

    def test_image_validation_without_arbitrary_count_or_size_caps(self):
        revision = 0
        for _ in range(5):
            revision = materials.add_image(self.project, "x.png", PNG, revision)["revision"]
        revision = materials.add_image(self.project, "six.png", PNG, revision)["revision"]
        with self.assertRaises(ValueError):
            materials.add_image(self.project, "bad.png", b"not an image", revision)
        source = self.root / "large.png"
        with source.open("wb") as output:
            output.write(PNG)
            output.seek(21 * 1024 * 1024)
            output.write(b"0")
        added = materials.add_image(self.project, "large.png", source, revision)
        revision = added["revision"]
        path, _ = materials.image_file(self.project, added["images"][-1]["id"])
        self.assertEqual(path.stat().st_size, source.stat().st_size)
        with self.assertRaises(ValueError):
            materials.edit_images(self.project, [{"id": "../fake", "note": ""}], revision)
        image = materials.get(self.project)["images"][0]
        with self.assertRaises(ValueError):
            materials.edit_images(self.project, [{"id": image["id"], "note": ""}] * 2, revision)

    def test_package_uses_current_view_and_includes_all_artifacts(self):
        job = storage.create_job(self.project, "source.wav", self.root / "source.wav", {})
        storage.update_job(job["id"], status="completed")
        materials.save_text(self.project, "summary", "Резюме", 0)
        materials.save_text(self.project, "minutes", "Протокол", 1)
        data = materials.add_image(self.project, "image.png", PNG, 2)
        materials.edit_images(self.project, [{"id": data["images"][0]["id"], "note": "Вариант"}], 3)
        with patch.object(materials.exporter, "combined_text", return_value="[00:01] Исходник") as combined:
            files = materials.package_files(self.project, True, "txt", "words", "plain", "on", "job:speaker")
        combined.assert_called_once_with(self.project, "words", "plain", "on", {"job:speaker"})
        archive = zipfile.ZipFile(io.BytesIO(materials.zip_files(files)))
        self.assertCountEqual(archive.namelist(), ["Резюме.docx", "Протокол.docx", "Схема 1.png", "Примечания к схемам.txt", "Транскрибация.txt"])
        self.assertEqual(archive.read("Схема 1.png"), PNG)
        self.assertEqual(archive.read("Транскрибация.txt").decode(), "[00:01] Исходник")

    def test_one_material_package_can_exclude_transcript(self):
        materials.save_text(self.project, "summary", "Итог", 0)
        with patch.object(materials.exporter, "combined_text") as combined:
            files = materials.package_files(self.project, False, "txt", "none", "formatted", "off")
        combined.assert_not_called()
        self.assertEqual([name for name, _ in files], ["Резюме.docx"])

    def test_missing_project_never_creates_directories(self):
        with self.assertRaises(KeyError):
            materials.save_text("missing", "summary", "x", 0)
        self.assertFalse((self.root / "sessions" / "missing").exists())

    def test_materials_follow_project_rename(self):
        materials.save_text(self.project, "summary", "Сохраняется", 0)
        storage.update_session(self.project, name="Другое имя")
        self.assertEqual(materials.get(self.project)["summary"], "Сохраняется")
        self.assertTrue((storage.session_dir(self.project) / "materials" / "Резюме.docx").exists())

    def test_reveal_opens_the_current_material_directory(self):
        materials.save_text(self.project, "summary", "Материал", 0)
        storage.update_session(self.project, name="Переименован")
        with patch("backend.app.os.startfile", create=True) as launch:
            self.assertEqual(materials_api.reveal(self.project), {"ok": True})
        launch.assert_called_once_with(str(storage.session_dir(self.project) / "materials"))

    def test_bulk_folder_export_preserves_hidden_projects_and_reports_errors(self):
        materials.save_text(self.project, "summary", "Результат", 0)
        storage.update_folder(self.folder["id"], feed_visible=False)
        empty = storage.create_session("Без материалов", self.folder["id"])["id"]
        result = materials_api.batch_export(materials_api.BatchExport(folder_id=self.folder["id"], include_transcript=False))
        self.assertEqual([item["project_id"] for item in result["projects"]], [self.project])
        self.assertEqual(result["errors"][0]["project_id"], empty)
        files = result["projects"][0]["files"]
        self.assertTrue(any(path.endswith("Резюме.docx") for path in files))
        self.assertTrue(any(path.endswith("Материалы.zip") for path in files))
        with zipfile.ZipFile(Path(result["directory"]) / "Пакет.zip") as archive:
            self.assertTrue(any(name.endswith("Резюме.docx") for name in archive.namelist()))
        states = {item["project_id"]: item for item in materials_api.status(self.folder["id"])["projects"]}
        self.assertTrue(states[self.project]["summary_ready"])
        self.assertFalse(states[empty]["summary_ready"])

    def test_prompt_api_and_catalog_share_exact_ui_text(self):
        with patch.object(materials_api.exporter, "combined_text", return_value="  Текст\n"):
            result = materials_api.prompt_input(self.project, "summary")
        source = json.loads((Path(__file__).resolve().parents[1] / "app" / "transcript-prompts.json").read_text(encoding="utf-8"))
        self.assertEqual(result["input"], source[0]["text"] + "\n\nТранскрибация:\n  Текст\n")
        self.assertEqual(materials_api.prompt_input(self.project, "summary", False)["input"], source[0]["text"])

    def test_diagram_prompt_adds_only_existing_materials_before_transcript(self):
        materials.save_text(self.project, "summary", "Резюме\n", 0)
        materials.save_text(self.project, "minutes", "Протокол", 1)
        with patch.object(materials_api.exporter, "combined_text", return_value="  Текущий текст\n"):
            result = materials_api.prompt_input(self.project, "diagram", include_materials=True)
            self.assertEqual(result["input"], "\n\n---\n\n".join([result["prompt"], "Резюме:\nРезюме\n", "Цифровой протокол:\nПротокол", "Исходная транскрибация:\n  Текущий текст\n"]))
            materials.save_text(self.project, "summary", "", 2)
            result = materials_api.prompt_input(self.project, "diagram", include_materials=True)
            self.assertNotIn("Резюме:\n", result["input"])
            self.assertIn("Цифровой протокол:\nПротокол", result["input"])
            unchanged = materials_api.prompt_input(self.project, "diagram")
            self.assertEqual(unchanged["input"], unchanged["prompt"] + "\n\nТранскрибация:\n  Текущий текст\n")

    def test_agent_can_export_fifty_projects_with_files_and_zips(self):
        ids = [self.project] + [storage.create_session(f"Пакет {i}", self.folder["id"])["id"] for i in range(49)]
        for project_id in ids:
            materials.save_text(project_id, "summary", f"Результат для {project_id}", 0)
        result = materials_api.batch_export(materials_api.BatchExport(project_ids=ids, include_transcript=False))
        self.assertFalse(result["errors"])
        self.assertEqual(len(result["projects"]), 50)
        for project in result["projects"]:
            self.assertEqual(len(project["files"]), 2)
            self.assertTrue(all(Path(path).is_file() for path in project["files"]))
        with zipfile.ZipFile(Path(result["directory"]) / "Пакет.zip") as archive:
            self.assertEqual(sum(name.endswith("Резюме.docx") for name in archive.namelist()), 50)
            self.assertEqual(sum(name.endswith("Материалы.zip") for name in archive.namelist()), 50)

    def test_api_conflict_is_409(self):
        materials_api.write(self.project, "minutes", materials_api.TextEdit(text="Ответ", revision=0))
        with self.assertRaises(HTTPException) as error:
            materials_api.write(self.project, "minutes", materials_api.TextEdit(text="Старое", revision=0))
        self.assertEqual(error.exception.status_code, 409)

    def test_failed_export_does_not_commit_new_text(self):
        materials.save_text(self.project, "summary", "Исходный текст", 0)
        with patch.object(materials, "_docx", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                materials.save_text(self.project, "summary", "Новый", 1)
        self.assertEqual(materials.get(self.project)["summary"], "Исходный текст")
        self.assertEqual((storage.session_dir(self.project) / "materials" / "Резюме.txt").read_text(encoding="utf-8"), "Исходный текст")

    def test_manifest_commit_failure_restores_derived_files(self):
        materials.save_text(self.project, "summary", "До изменения", 0)
        original = materials.atomic_write
        def fail_manifest(path, content):
            if path.name == "materials.json":
                raise OSError("Manifest disk error")
            return original(path, content)
        with patch.object(materials, "atomic_write", side_effect=fail_manifest):
            with self.assertRaises(OSError):
                materials.save_text(self.project, "summary", "Не записалось", 1)
        root = storage.session_dir(self.project) / "materials"
        self.assertEqual(materials.get(self.project)["summary"], "До изменения")
        self.assertEqual((root / "Резюме.txt").read_text(encoding="utf-8"), "До изменения")
        self.assertIn("До изменения", [paragraph.text for paragraph in Document(root / "Резюме.docx").paragraphs])
        self.assertFalse((root / ".exports-dirty").exists())

    def test_interrupted_export_recovers_from_committed_manifest(self):
        materials.save_text(self.project, "summary", "Верный результат", 0)
        root = storage.session_dir(self.project) / "materials"
        (root / "Резюме.txt").write_text("Не завершённая запись", encoding="utf-8")
        (root / ".exports-dirty").write_bytes(b"1")
        self.assertEqual(materials.get(self.project)["summary"], "Верный результат")
        self.assertEqual((root / "Резюме.txt").read_text(encoding="utf-8"), "Верный результат")

    def test_idempotent_text_retry_preserves_other_materials(self):
        materials.save_text(self.project, "summary", "Итог", 0)
        materials.save_text(self.project, "minutes", "Другая запись", 1)
        result = materials.save_text(self.project, "summary", "Итог", 0)
        self.assertEqual(result["revision"], 2)
        self.assertEqual(result["minutes"], "Другая запись")

    def test_corrupt_manifest_is_not_replaced_by_empty_data(self):
        materials.save_text(self.project, "summary", "Важный текст", 0)
        path = storage.session_dir(self.project) / "materials" / "materials.json"
        path.write_text('{"broken":', encoding="utf-8")
        with self.assertRaises(materials.InvalidState):
            materials.save_text(self.project, "summary", "Замена", 1)
        self.assertEqual(path.read_text(encoding="utf-8"), '{"broken":')

    def test_directory_rename_waits_for_material_save(self):
        started, proceed = threading.Event(), threading.Event()
        original = materials.atomic_write
        def pause(path, content):
            if path.name == "Резюме.txt":
                started.set()
                self.assertTrue(proceed.wait(10))
            return original(path, content)
        with ThreadPoolExecutor(max_workers=2) as pool, patch.object(materials, "atomic_write", side_effect=pause):
            save = pool.submit(materials.save_text, self.project, "summary", "При переименовании", 0)
            self.assertTrue(started.wait(10))
            rename = pool.submit(storage.update_session, self.project, name="Новое имя")
            self.assertFalse(rename.done())
            proceed.set(); save.result(10); rename.result(10)
        root = storage.session_dir(self.project) / "materials"
        self.assertEqual(materials.get(self.project)["summary"], "При переименовании")
        self.assertEqual((root / "Резюме.txt").read_text(encoding="utf-8"), "При переименовании")

    def test_failed_batch_project_never_leaks_partial_files_into_archive(self):
        materials.save_text(self.project, "summary", "Результат", 0)
        other = storage.create_session("Успешный проект")["id"]
        materials.save_text(other, "minutes", "Протокол", 0)
        original = materials.atomic_write
        def fail_one(path, content):
            if path.name == "Резюме.docx" and ".stage-" in str(path):
                raise OSError("Export failed")
            return original(path, content)
        with patch.object(materials, "atomic_write", side_effect=fail_one):
            result = material_exports.export_projects([self.project, other], "both", False, "txt", "none", "formatted", "off")
        self.assertEqual([project["project_id"] for project in result["projects"]], [other])
        self.assertEqual([project["project_id"] for project in result["errors"]], [self.project])
        with zipfile.ZipFile(Path(result["directory"]) / "Пакет.zip") as archive:
            self.assertFalse(any(name.endswith("Резюме.docx") or ".stage-" in name for name in archive.namelist()))
            self.assertTrue(any(name.endswith("Протокол.docx") for name in archive.namelist()))

    def test_material_only_project_can_download_default_package(self):
        materials.save_text(self.project, "summary", "Только резюме", 0)
        self.assertEqual([name for name, _ in materials.package_files(self.project, True, "txt", "none", "formatted", "off")], ["Резюме.docx"])


if __name__ == "__main__":
    unittest.main()
