from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any

from docx import Document

from . import hallucinations, paragraphs, storage

logger = logging.getLogger("whisper.exporter")

SpeakerAliases = dict[tuple[str, str], str]
SpeakerSelection = set[str] | None

# Готовые файлы проекта, которые сервис держит на диске сам. Имя общее для всех
# трёх, чтобы внешний скрипт находил расшифровку по маске, не зная формата.
AUTO_EXPORT_KINDS = ("md", "txt", "docx")
AUTO_EXPORT_SUFFIX = " — расшифровка"


def _clock(seconds: float, decimal: str = ",") -> str:
    milliseconds = max(0, round(seconds * 1000))
    hours, remainder = divmod(milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    secs, millis = divmod(remainder, 1000)
    return f"{hours:02}:{minutes:02}:{secs:02}{decimal}{millis:03}"


def _project_file_name(project: dict[str, Any]) -> str:
    """Имя файла проекта — название проекта, пригодное для Windows."""
    return storage.safe_file_name(project["name"], fallback="Проект")


def _plain_text(text: str) -> str:
    text = text.lower().replace("ё", "е")
    text = re.sub(r"[^\w\s-]", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text).strip()


def _display_text(text: str, formatting: str) -> str:
    clean = text.strip()
    return _plain_text(clean) if formatting == "plain" else clean


def _timed_words(segment: dict[str, Any]) -> list[dict[str, Any]]:
    existing = segment.get("words") or []
    if existing:
        return [
            {
                "start": float(word.get("start") or segment.get("start") or 0),
                "end": float(word.get("end") or word.get("start") or 0),
                "text": str(word.get("word") or "").strip(),
                "speaker": word.get("speaker") or segment.get("speaker"),
            }
            for word in existing
            if str(word.get("word") or "").strip()
        ]
    # Старые результаты не содержат word timestamps. Для них даём
    # равномерную оценку внутри точной временной границы фразы.
    words = re.findall(r"\S+", str(segment.get("text") or ""))
    if not words:
        return []
    start = float(segment.get("start") or 0)
    end = max(start, float(segment.get("end") or start))
    step = (end - start) / len(words) if end > start else 0
    return [
        {
            "start": start + index * step,
            "end": start + (index + 1) * step,
            "text": word,
            "speaker": segment.get("speaker"),
        }
        for index, word in enumerate(words)
    ]


def _speaker_key(job_id: str, value: Any) -> str:
    return f"{job_id}:{value}"


def _speaker_label(
    value: Any,
    job_id: str = "",
    aliases: SpeakerAliases | None = None,
) -> str:
    if value is None:
        return ""
    raw = str(value)
    if aliases and job_id and (job_id, raw) in aliases:
        return aliases[(job_id, raw)]
    match = re.search(r"(\d+)$", raw)
    if match:
        return f"СПИКЕР {int(match.group(1)) + 1}"
    return raw.replace("_", " ").upper()


def _speaker_selected(
    job_id: str,
    value: Any,
    selected: SpeakerSelection,
) -> bool:
    if selected is None:
        return True
    if value is None:
        return False
    return _speaker_key(job_id, value) in selected


def parse_speaker_keys(value: str | None) -> SpeakerSelection:
    if value is None or value == "" or value == "all":
        return None
    if value == "none":
        return set()
    return {item for item in value.split(",") if item}


def _filtered_segments(
    segments: list[dict[str, Any]],
    job_id: str,
    selected: SpeakerSelection,
) -> list[dict[str, Any]]:
    if selected is None:
        return segments
    filtered: list[dict[str, Any]] = []
    for segment in segments:
        words = segment.get("words") or []
        if words:
            kept_words = [
                dict(word)
                for word in words
                if _speaker_selected(
                    job_id,
                    word.get("speaker") or segment.get("speaker"),
                    selected,
                )
            ]
            if not kept_words:
                continue
            item = dict(segment)
            item["words"] = kept_words
            item["start"] = float(
                kept_words[0].get("start") or segment.get("start") or 0
            )
            item["end"] = float(
                kept_words[-1].get("end") or segment.get("end") or item["start"]
            )
            item["text"] = " ".join(
                str(word.get("word") or "").strip()
                for word in kept_words
                if str(word.get("word") or "").strip()
            )
            speakers = [
                word.get("speaker") or segment.get("speaker")
                for word in kept_words
            ]
            item["speaker"] = next(
                (speaker for speaker in speakers if speaker is not None),
                segment.get("speaker"),
            )
            filtered.append(item)
        elif _speaker_selected(job_id, segment.get("speaker"), selected):
            filtered.append(segment)
    return filtered


def _speaker_heading(label: str, markup: str) -> str:
    return f"**{label}:**" if markup == "markdown" else f"{label}:"


def _render_file_text(
    segments: list[dict[str, Any]],
    timestamps: str,
    formatting: str,
    speakers: str,
    job_id: str,
    aliases: SpeakerAliases,
    selected: SpeakerSelection,
    markup: str = "text",
) -> str:
    segments = _filtered_segments(segments, job_id, selected)
    lines: list[str] = []
    current_speaker = ""
    if timestamps == "words":
        for segment in segments:
            for word in _timed_words(segment):
                text = _display_text(word["text"], formatting)
                if not text:
                    continue
                label = (
                    _speaker_label(word.get("speaker"), job_id, aliases)
                    if speakers == "on"
                    else ""
                )
                prefix = f"[{_clock(word['start'], '.')}] "
                if label and label != current_speaker:
                    lines.append(f"\n{label}:")
                    current_speaker = label
                lines.append(f"{prefix}{text}")
        return "\n".join(lines).strip()

    # Абзацы считаются по исходному тексту модели: режим «Без знаков» снимает
    # точки, по которым граница и определяется, поэтому очистка идёт после.
    blocks: list[str] = []
    for group in paragraphs.group(segments):
        pieces = [
            (
                item,
                _display_text(str(item.get("text") or ""), formatting),
            )
            for item in group
        ]
        pieces = [(item, text) for item, text in pieces if text]
        if not pieces:
            continue
        if timestamps == "segments":
            body = "\n".join(
                f"[{_clock(float(item['start']), '.').split('.')[0]}] {text}"
                for item, text in pieces
            )
        else:
            body = " ".join(text for _, text in pieces)
        label = (
            _speaker_label(
                paragraphs.segment_speaker(group[0]),
                job_id,
                aliases,
            )
            if speakers == "on"
            else ""
        )
        if label and label != current_speaker:
            blocks.append(f"{_speaker_heading(label, markup)}\n{body}")
            current_speaker = label
        else:
            blocks.append(body)
    return "\n\n".join(blocks).strip()


def render_segments(
    segments: list[dict[str, Any]],
    timestamps: str = "none",
    formatting: str = "formatted",
    speakers: str = "off",
    job_id: str = "",
    aliases: SpeakerAliases | None = None,
    markup: str = "text",
) -> str:
    """Текст одного файла — тем же кодом, что и просмотр с экспортом.

    Нужна распознавателю: он пишет свои txt, md и docx сразу после работы, и
    собирать абзацы вторым способом означало бы две разные расшифровки одной
    записи.
    """
    return _render_file_text(
        segments,
        timestamps,
        formatting,
        speakers,
        job_id,
        aliases or {},
        None,
        markup,
    )


def _completed_project(project_id: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    project = storage.get_session(project_id)
    jobs = [
        job
        for job in storage.list_jobs(session_id=project_id)
        if job["status"] == "completed"
    ]
    if not jobs:
        raise RuntimeError("В проекте пока нет готовых расшифровок")
    return project, jobs


def _job_text(job: dict[str, Any]) -> str:
    path_value = job["outputs"].get("txt")
    if not path_value:
        return ""
    path = Path(path_value)
    return path.read_text(encoding="utf-8").strip() if path.exists() else ""


def combined_text(
    project_id: str,
    timestamps: str = "none",
    formatting: str = "formatted",
    speakers: str = "off",
    speaker_keys: SpeakerSelection = None,
    markup: str = "text",
) -> str:
    _, jobs, files = _project_payload(project_id)
    by_job = {item["job_id"]: item for item in files}
    aliases = storage.list_speaker_aliases(project_id)
    parts: list[str] = []
    for job in jobs:
        payload = by_job.get(job["id"])
        if payload:
            text = _render_file_text(
                payload["segments"],
                timestamps,
                formatting,
                speakers,
                job["id"],
                aliases,
                speaker_keys,
                markup,
            )
        else:
            # Совместимость со старыми заданиями, где JSON мог отсутствовать.
            if speaker_keys is not None:
                continue
            text = _job_text(job)
            if formatting == "plain":
                text = _plain_text(text)
        if not text.strip():
            continue
        if len(jobs) > 1:
            parts.append(f"## {job['original_name']}\n\n{text.strip()}")
        else:
            parts.append(text.strip())
    return "\n\n".join(parts).strip()


def write_docx(title: str, text: str, target: Path) -> Path:
    """Word из готового текста: пустые строки в документе не нужны.

    Абзацы отделяются друг от друга пустой строкой, но в Word отступ между
    абзацами задаёт сам стиль. Если переносить разделители как есть, документ
    получает вдвое больше пустых абзацев, чем содержательных.
    """
    document = Document()
    document.add_heading(title, level=0)
    for line in text.splitlines():
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.strip():
            document.add_paragraph(line)
    document.save(target)
    return target


def create_job_docx(job: dict[str, Any]) -> Path:
    text = _job_text(job)
    if not text:
        raise RuntimeError("Текст ещё не готов")
    output_dir = storage.session_dir(job["session_id"]) / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = storage.job_file_base(job)
    target = output_dir / f"{stem}.docx"
    write_docx(stem, text, target)
    storage.add_output(job["id"], "docx", target)
    return target


def refresh_project_files(project_id: str) -> list[Path]:
    """Обновить готовые файлы проекта на диске.

    Раньше результат жил только в базе: чтобы получить файл, надо было открыть
    проект и нажать «Скачать». Если Whisper — стадия конвейера, а работа идёт
    дальше в других программах, это лишний заход в интерфейс на каждую запись.

    Поэтому рядом с проектом всегда лежат три файла с одинаковым именем:
    `.md` для передачи в модель, `.txt` для чего угодно и `.docx` для чтения
    и правки. Обновляются они сами — после каждого распознавания, разметки
    спикеров и переименования спикера, — поэтому в папке никогда не лежит
    вчерашняя версия.

    Вид берётся из настроек просмотра: таймкоды и пунктуация — как выбрано,
    спикеры — включены всегда, когда разметка есть. Настройка `speakers`
    управляет переключателем в просмотрщике, который можно щёлкать туда-сюда, а
    файл на диске должен нести всё, что о записи известно.

    Ошибка записи не считается ошибкой распознавания: открытый в Word файл
    заблокирован Windows, и попытка его перезаписать не должна ронять задание.
    Что не удалось записать — просто не попадает в ответ.
    """
    try:
        return _write_project_files(project_id)
    except Exception:
        # Обновление вызывается сразу после того, как задание объявлено
        # готовым, и стоит внутри общего `except` распознавателя. Без этой
        # заглушки любая неожиданность здесь пометила бы шестичасовую
        # расшифровку неудачной, хотя она уже лежит в базе целой.
        logger.exception(
            "Не удалось обновить готовые файлы проекта %s", project_id
        )
        return []


def _write_project_files(project_id: str) -> list[Path]:
    try:
        project, _, files = _project_payload(project_id)
    except (RuntimeError, KeyError, OSError, json.JSONDecodeError):
        return []
    settings = storage.get_settings()
    speakers = (
        "on"
        if any(
            paragraphs.segment_speaker(segment) is not None
            for file in files
            for segment in file["segments"]
        )
        else "off"
    )
    rendered = {
        markup: combined_text(
            project_id,
            settings.get("default_view_timestamps", "none"),
            settings.get("default_view_formatting", "formatted"),
            speakers,
            None,
            markup,
        )
        for markup in ("text", "markdown")
    }
    directory = storage.session_dir(project_id)
    directory.mkdir(parents=True, exist_ok=True)
    base = f"{_project_file_name(project)}{AUTO_EXPORT_SUFFIX}"
    written: list[Path] = []
    for kind in AUTO_EXPORT_KINDS:
        target = directory / f"{base}.{kind}"
        try:
            if kind == "md":
                target.write_text(
                    f"# {project['name']}\n\n{rendered['markdown']}\n",
                    encoding="utf-8",
                )
            elif kind == "txt":
                target.write_text(rendered["text"] + "\n", encoding="utf-8")
            else:
                write_docx(project["name"], rendered["text"], target)
        except OSError:
            continue
        written.append(target)
    _drop_stale_project_files(directory, {item.name for item in written})
    return written


def backfill_project_files() -> int:
    """Дописать готовые файлы проектам, у которых их ещё нет.

    Нужно один раз — проектам, законченным до того, как файлы стали
    появляться сами. Проверка дешёвая: если все три файла на месте, проект
    пропускается, ничего не читая и не перезаписывая.
    """
    with storage.connect() as db:
        ids = [
            row["id"]
            for row in db.execute("SELECT id FROM sessions").fetchall()
        ]
    written = 0
    for project_id in ids:
        try:
            project = storage.get_session(project_id)
        except (KeyError, OSError):
            continue
        base = f"{_project_file_name(project)}{AUTO_EXPORT_SUFFIX}"
        directory = storage.session_dir(project_id)
        if all(
            (directory / f"{base}.{kind}").exists()
            for kind in AUTO_EXPORT_KINDS
        ):
            continue
        if refresh_project_files(project_id):
            written += 1
    return written


def _drop_stale_project_files(directory: Path, keep: set[str]) -> None:
    """Убрать файлы, оставшиеся от прежнего названия проекта.

    Имя файла повторяет название проекта, поэтому после переименования рядом
    оказались бы две расшифровки — свежая и брошенная. Под раздачу попадают
    только файлы с нашим суффиксом: всё остальное в папке принесено человеком.
    """
    try:
        entries = list(directory.iterdir())
    except OSError:
        return
    for item in entries:
        if (
            item.name not in keep
            and item.is_file()
            and item.suffix[1:].lower() in AUTO_EXPORT_KINDS
            and item.stem.endswith(AUTO_EXPORT_SUFFIX)
        ):
            try:
                item.unlink()
            except OSError:
                continue


def _project_payload(
    project_id: str,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    project, jobs = _completed_project(project_id)
    files: list[dict[str, Any]] = []
    for job in jobs:
        json_value = job["outputs"].get("json")
        if not json_value or not Path(json_value).exists():
            continue
        payload = json.loads(Path(json_value).read_text(encoding="utf-8"))
        options = payload.get("options") or {}
        files.append(
            {
                "job_id": job["id"],
                "source": job["original_name"],
                "duration_seconds": float(payload.get("duration_seconds") or 0),
                # Метки обрезанного файла живут на шкале исходника. Для сплошной
                # ленты субтитров их надо вернуть к нулю; отдельные форматы
                # (txt, docx, json) показывают время исходника как есть.
                "time_offset": float(options.get("trim_start_seconds") or 0),
                "language": payload.get("language"),
                "options": options,
                # Фильтр здесь — сетка безопасности для результатов, записанных
                # до его появления: просмотр и все экспорты проекта идут через
                # это место, поэтому старая расшифровка выходит чистой без
                # повторного распознавания. Для новых результатов JSON уже
                # чистый, и вызов ничего не меняет.
                "segments": hallucinations.clean_segments(
                    payload.get("segments", [])
                )[0],
            }
        )
    return project, jobs, files


def project_speakers(project_id: str) -> list[dict[str, Any]]:
    _, jobs, files = _project_payload(project_id)
    aliases = storage.list_speaker_aliases(project_id)
    by_job = {item["job_id"]: item for item in files}
    result: list[dict[str, Any]] = []
    for job in jobs:
        payload = by_job.get(job["id"])
        if not payload:
            continue
        counts: dict[str, int] = {}
        durations: dict[str, float] = {}
        for segment in payload["segments"]:
            words = segment.get("words") or []
            speaker_values: set[str] = set()
            if words:
                for word in words:
                    speaker = word.get("speaker") or segment.get("speaker")
                    if speaker is None:
                        continue
                    speaker_id = str(speaker)
                    speaker_values.add(speaker_id)
                    durations[speaker_id] = durations.get(
                        speaker_id, 0.0
                    ) + max(
                        0.0,
                        float(word.get("end") or 0)
                        - float(word.get("start") or 0),
                    )
            if not speaker_values and segment.get("speaker"):
                speaker_values = {str(segment["speaker"])}
                speaker_id = next(iter(speaker_values))
                durations[speaker_id] = durations.get(
                    speaker_id, 0.0
                ) + max(
                    0.0,
                    float(segment.get("end") or 0)
                    - float(segment.get("start") or 0),
                )
            for speaker_id in speaker_values:
                counts[speaker_id] = counts.get(speaker_id, 0) + 1
        for speaker_id in sorted(counts, key=_speaker_sort_key):
            fallback = _speaker_label(speaker_id)
            custom_name = aliases.get((job["id"], speaker_id), "")
            result.append(
                {
                    "key": _speaker_key(job["id"], speaker_id),
                    "job_id": job["id"],
                    "speaker_id": speaker_id,
                    "source": job["original_name"],
                    "default_name": fallback,
                    "custom_name": custom_name,
                    "name": custom_name or fallback,
                    "segment_count": counts[speaker_id],
                    "duration_seconds": round(durations[speaker_id], 3),
                }
            )
    return result


def _speaker_sort_key(value: str) -> tuple[int, str]:
    match = re.search(r"(\d+)$", value)
    return (int(match.group(1)) if match else 1_000_000, value)


def create_project_export(
    project_id: str,
    kind: str,
    timestamps: str = "none",
    formatting: str = "formatted",
    speakers: str = "off",
    speaker_keys: SpeakerSelection = None,
) -> Path:
    if kind not in {"md", "txt", "docx", "srt", "vtt", "json"}:
        raise ValueError(kind)
    project, jobs, files = _project_payload(project_id)
    aliases = storage.list_speaker_aliases(project_id)
    export_dir = storage.session_dir(project_id) / "project_exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    timestamp_name = {
        "none": "без-таймкодов",
        "segments": "таймкоды-по-фразам",
        "words": "таймкоды-по-словам",
    }.get(timestamps, "без-таймкодов")
    variant = (
        timestamp_name
        + "-"
        + ("без-пунктуации" if formatting == "plain" else "с-пунктуацией")
        + ("-по-спикерам" if speakers == "on" else "")
    )
    target = export_dir / f"{_project_file_name(project)} - {variant}.{kind}"

    if kind in {"txt", "md", "docx"}:
        text = combined_text(
            project_id,
            timestamps,
            formatting,
            speakers,
            speaker_keys,
            markup="markdown" if kind == "md" else "text",
        )
        if kind == "docx":
            write_docx(project["name"], text, target)
        elif kind == "md":
            target.write_text(
                f"# {project['name']}\n\n{text}\n", encoding="utf-8"
            )
        else:
            target.write_text(text + "\n", encoding="utf-8")
        return target

    if kind == "json":
        target.write_text(
            json.dumps(
                {
                    "project": {
                        "id": project["id"],
                        "name": project["name"],
                        "folder": project.get("folder_name"),
                    },
                    "view": {
                        "timestamps": timestamps,
                        "formatting": formatting,
                        "speakers": speakers,
                        "speaker_keys": (
                            sorted(speaker_keys)
                            if speaker_keys is not None
                            else "all"
                        ),
                    },
                    "speaker_names": [
                        {
                            "key": item["key"],
                            "source": item["source"],
                            "speaker_id": item["speaker_id"],
                            "name": item["name"],
                        }
                        for item in project_speakers(project_id)
                    ],
                    "files": [
                        {
                            **file,
                            "segments": [
                                {
                                    **segment,
                                    "text": _display_text(
                                        str(segment.get("text") or ""),
                                        formatting,
                                    ),
                                    "speaker_name": _speaker_label(
                                        segment.get("speaker"),
                                        file["job_id"],
                                        aliases,
                                    ),
                                    "words": [
                                        {
                                            **word,
                                            "word": _display_text(
                                                str(word.get("word") or ""),
                                                formatting,
                                            ),
                                            "speaker_name": _speaker_label(
                                                word.get("speaker")
                                                or segment.get("speaker"),
                                                file["job_id"],
                                                aliases,
                                            ),
                                        }
                                        for word in (segment.get("words") or [])
                                    ],
                                }
                                for segment in _filtered_segments(
                                    file["segments"],
                                    file["job_id"],
                                    speaker_keys,
                                )
                            ],
                        }
                        for file in files
                    ],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return target

    offset = 0.0
    blocks: list[str] = ["WEBVTT"] if kind == "vtt" else []
    index = 1
    for file in files:
        filtered_segments = _filtered_segments(
            file["segments"],
            file["job_id"],
            speaker_keys,
        )
        timed_items = (
            [
                word
                for segment in filtered_segments
                for word in _timed_words(segment)
            ]
            if timestamps == "words"
            else [
                {
                    "start": float(segment["start"]),
                    "end": float(segment["end"]),
                    "text": str(segment["text"]),
                    "speaker": segment.get("speaker"),
                }
                for segment in filtered_segments
            ]
        )
        local = offset - float(file["time_offset"])
        for item in timed_items:
            start = local + float(item["start"])
            end = local + float(item["end"])
            label = (
                _speaker_label(
                    item.get("speaker"),
                    file["job_id"],
                    aliases,
                )
                if speakers == "on"
                else ""
            )
            rendered = _display_text(str(item["text"]), formatting)
            if label:
                rendered = f"{label}: {rendered}"
            if kind == "srt":
                blocks.append(
                    f"{index}\n{_clock(start)} --> {_clock(end)}\n"
                    f"{rendered}"
                )
            else:
                blocks.append(
                    f"{_clock(start, '.')} --> {_clock(end, '.')}\n"
                    f"{rendered}"
                )
            index += 1
        offset += file["duration_seconds"]
    target.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
    return target
