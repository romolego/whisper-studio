"""Project artifacts, independent of inference and HTTP.

The manifest is the source of truth. Revisions prevent a stale browser/agent
from replacing another editor's work. Files stay inside the project directory;
an adapter for a corporate object/attachment store can replace this module.
"""
from __future__ import annotations

import io
import json
import logging
import os
import re
import shutil
import uuid
import zipfile
from pathlib import Path

import av
from docx import Document

from . import exporter, storage
from .project_locks import project_lock

TEXT_KINDS = {"summary": "Резюме", "minutes": "Протокол"}
logger = logging.getLogger(__name__)


class Conflict(ValueError):
    pass


class InvalidState(RuntimeError):
    pass


def _root(project_id: str) -> Path:
    storage.get_session(project_id)  # Never create files for a nonexistent ID.
    return storage.session_dir(project_id) / "materials"


def _load(root: Path) -> dict:
    path = root / "materials.json"
    try:
        data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {
            "schema_version": 1, "revision": 0, "updated_at": None, "summary": "", "minutes": "", "images": []}
        valid = (isinstance(data, dict) and type(data.get("schema_version")) is int and data["schema_version"] == 1
                 and type(data.get("revision")) is int and data["revision"] >= 0
                 and all(isinstance(data.get(kind), str) and len(data[kind]) <= 2_000_000 for kind in TEXT_KINDS)
                 and isinstance(data.get("images"), list))
        if not valid:
            raise ValueError("Invalid manifest")
        ids = set()
        for image in data["images"]:
            if (not isinstance(image, dict) or not re.fullmatch(r"[a-f0-9]{32}", str(image.get("id", "")))
                or image["id"] in ids or image.get("file") not in {f"{image['id']}.{ext}" for ext in ("png", "jpg", "webp")}
                or not isinstance(image.get("name"), str) or not isinstance(image.get("note"), str) or len(image["note"]) > 10000):
                raise ValueError("Invalid image manifest")
            ids.add(image["id"])
    except (ValueError, TypeError) as exc:
        raise InvalidState("Файл материалов повреждён или имеет неизвестную версию. Он оставлен без изменений.") from exc
    if (root / ".exports-dirty").exists():
        _recover_text_files(root, data)
    return data


def atomic_write(path: Path, content: bytes | Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        with temp.open("wb") as file:
            if isinstance(content, Path):
                with content.open("rb") as source:
                    shutil.copyfileobj(source, file, 1024 * 1024)
            else:
                file.write(content)
            file.flush()
            os.fsync(file.fileno())
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def _commit(root: Path, data: dict) -> dict:
    data["revision"] += 1
    data["updated_at"] = storage.utc_now()
    atomic_write(root / "materials.json", json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"))
    return data


def _check(data: dict, revision: int) -> None:
    if revision != data["revision"]:
        raise Conflict("Материалы изменены в другом окне. Обновите их перед сохранением.")


def get(project_id: str) -> dict:
    with project_lock(project_id):
        return _load(_root(project_id))


def _docx(text: str, title: str | None = None) -> bytes:
    document = Document()
    if title:
        document.add_heading(title, 0)
    for line in text.splitlines():
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        bullet = re.match(r"^\s*[-*•]\s+(.*)", line)
        paragraph = (document.add_heading(level=min(3, len(heading[1]))) if heading
                     else document.add_paragraph(style="List Bullet" if bullet else None))
        value = heading[2] if heading else bullet[1] if bullet else line
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        for part in re.split(r"(\*\*.*?\*\*)", value):
            run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
            if part.startswith("**") and part.endswith("**"):
                run.bold = True
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def save_text(project_id: str, kind: str, text: str, revision: int) -> dict:
    if kind not in TEXT_KINDS or len(text) > 2_000_000:
        raise ValueError("Неизвестный материал или слишком длинный текст")
    with project_lock(project_id):
        root = _root(project_id)
        data = _load(root)
        if data[kind] == text:
            return data  # Idempotent retry after a lost acknowledgement.
        _check(data, revision)
        data[kind] = text
        contents = _text_files(kind, text)  # Render before changing anything.
        marker = root / ".exports-dirty"
        atomic_write(marker, b"1")
        try:
            _write_text_files(root, contents)
            result = _commit(root, data)
        except Exception:
            # Manifest commit is atomic. Restore derived files from the last
            # committed state; keep the marker if the disk is still unavailable.
            try:
                _load(root)
            except (OSError, InvalidState):
                logger.warning("Artifact exports need recovery", exc_info=True)
            raise
        _remove_generated(marker)
        return result


def _text_files(kind: str, text: str) -> dict[str, bytes | None]:
    contents = {"txt": text.encode("utf-8"), "docx": _docx(text)} if text.strip() else {}
    return {f"{TEXT_KINDS[kind]}.{ext}": contents.get(ext) for ext in ("txt", "docx")}


def _write_text_files(root: Path, contents: dict[str, bytes | None]) -> None:
    for name, content in contents.items():
        target = root / name
        if content is None:
            target.unlink(missing_ok=True)
        else:
            atomic_write(target, content)


def _recover_text_files(root: Path, data: dict) -> None:
    contents = {}
    for kind in TEXT_KINDS:
        contents.update(_text_files(kind, data[kind]))
    _write_text_files(root, contents)
    _remove_generated(root / ".exports-dirty")


def _remove_generated(path: Path) -> None:
    try:
        path.unlink(missing_ok=True)
    except OSError:
        logger.warning("Could not remove generated artifact file: %s", path)


def _image_extension(content: bytes | Path) -> str:
    if isinstance(content, Path):
        with content.open("rb") as source:
            header = source.read(16)
    else:
        header = content[:16]
    if header.startswith(b"\x89PNG\r\n\x1a\n"):
        extension = "png"
    elif header.startswith(b"\xff\xd8\xff"):
        extension = "jpg"
    elif header.startswith(b"RIFF") and header[8:12] == b"WEBP":
        extension = "webp"
    else:
        raise ValueError("Нужна картинка PNG, JPEG или WebP")
    try:
        with av.open(str(content) if isinstance(content, Path) else io.BytesIO(content)) as image:
            stream = image.streams.video[0]
            if stream.width <= 0 or stream.height <= 0:
                raise ValueError("Не удалось определить размер картинки")
            next(image.decode(stream))
    except Exception as exc:
        raise ValueError("Не удалось прочитать картинку (PNG, JPEG или WebP)") from exc
    return extension


def add_image(project_id: str, name: str, content: bytes | Path, revision: int) -> dict:
    if (isinstance(content, Path) and content.stat().st_size == 0) or (isinstance(content, bytes) and not content):
        raise ValueError("Файл картинки пуст")
    extension = _image_extension(content)
    with project_lock(project_id):
        root = _root(project_id)
        data = _load(root)
        _check(data, revision)
        image_id = uuid.uuid4().hex
        filename = f"{image_id}.{extension}"
        atomic_write(root / "images" / filename, content)
        clean_name = storage.safe_file_name(name, fallback="Схема")
        clean_name = str(Path(clean_name).with_suffix(f".{extension}"))
        data["images"].append({"id": image_id, "file": filename,
                               "name": clean_name,
                               "note": ""})
        try:
            return _commit(root, data)
        except Exception:
            (root / "images" / filename).unlink(missing_ok=True)
            raise


def edit_images(project_id: str, images: list[dict], revision: int) -> dict:
    with project_lock(project_id):
        root = _root(project_id)
        data = _load(root)
        if images == [{"id": image["id"], "note": image["note"]} for image in data["images"]]:
            return data
        _check(data, revision)
        existing = {image["id"]: image for image in data["images"]}
        ids = [image["id"] for image in images]
        if len(ids) != len(set(ids)) or any(key not in existing for key in ids):
            raise ValueError("Неизвестная или повторяющаяся схема")
        data["images"] = [{**existing[image["id"]], "note": image["note"]} for image in images]
        result = _commit(root, data)
        for key in existing.keys() - set(ids):
            _remove_generated(root / "images" / existing[key]["file"])
        return result


def image_file(project_id: str, image_id: str) -> tuple[Path, str]:
    with project_lock(project_id):
        root = _root(project_id)
        item = next((item for item in _load(root)["images"] if item["id"] == image_id), None)
        if item is None:
            raise KeyError(image_id)
        return root / "images" / item["file"], item["name"]


def image_bytes(project_id: str, image_id: str) -> tuple[bytes, str, str]:
    # Return a snapshot while holding the project lock, before a rename/delete.
    with project_lock(project_id):
        path, name = image_file(project_id, image_id)
        return path.read_bytes(), name, path.suffix


def text_file(project_id: str, kind: str, extension: str) -> tuple[str, bytes]:
    if kind not in TEXT_KINDS or extension not in {"txt", "docx"}:
        raise ValueError("Неизвестный формат материала")
    text = get(project_id)[kind]
    if not text.strip():
        raise ValueError("Материал ещё не заполнен")
    return (f"{TEXT_KINDS[kind]}.{extension}", text.encode("utf-8") if extension == "txt"
            else _docx(text))


def package_files(project_id: str, include_transcript: bool, transcript_format: str,
                  timestamps: str, formatting: str, speakers: str,
                  speaker_keys: str | None = None) -> list[tuple[str, bytes]]:
    if transcript_format not in {"txt", "docx"}:
        raise ValueError("Неизвестный формат транскрибации")
    with project_lock(project_id):
        data = get(project_id)
        files = [text_file(project_id, kind, "docx") for kind in TEXT_KINDS if data[kind].strip()]
        notes = []
        for index, image in enumerate(data["images"], 1):
            path, _ = image_file(project_id, image["id"])
            name = f"Схема {index}{path.suffix}"
            files.append((name, path.read_bytes()))
            if image["note"].strip():
                notes.append(f"{name}: {image['note']}")
        if notes:
            files.append(("Примечания к схемам.txt", "\n\n".join(notes).encode("utf-8")))
        if not files:
            raise ValueError("В проекте ещё нет сохранённых материалов")
        if include_transcript and storage.get_session(project_id)["completed_count"]:
            text = exporter.combined_text(project_id, timestamps, formatting, speakers,
                                          exporter.parse_speaker_keys(speaker_keys))
            if not text.strip():
                raise ValueError("Транскрибация в выбранном виде пуста. Измените фильтр или исключите её из пакета.")
            files.append((f"Транскрибация.{transcript_format}", text.encode("utf-8")
                          if transcript_format == "txt" else _docx(text, title="Транскрибация")))
        return files


def zip_files(files: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files:
            archive.writestr(name, content)
    return output.getvalue()


def directory(project_id: str) -> Path:
    with project_lock(project_id):
        root = _root(project_id)
        if not root.is_dir():
            raise KeyError(project_id)
        return root
